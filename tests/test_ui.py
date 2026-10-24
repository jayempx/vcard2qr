"""Exercise Qt state, export actions, worker completion and cooperative shutdown."""

import time

from openpyxl import Workbook
from PySide6.QtCore import Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QFileDialog

from vcard2qr.app import VCardQRApp


def wait_for_job(qapp, window):
    """Pump queued signals with a deadline rather than sleeping the GUI thread."""
    deadline = time.monotonic() + 15
    while window.worker is not None and time.monotonic() < deadline:
        qapp.processEvents()
        QTest.qWait(5)
    assert window.worker is None


def test_generate_edit_copy_save_and_field_error(qapp, tmp_path, monkeypatch):
    window = VCardQRApp()
    window.show()
    window.fields["First Name"].setText("Ada")
    QTest.mouseClick(window.generate_button, Qt.MouseButton.LeftButton)
    assert window.save_button.isEnabled()
    window.copy_to_clipboard()
    assert qapp.clipboard().image().size().width() == 512
    output = tmp_path / "saved.png"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *args: (str(output), ""))
    window.save_current()
    assert output.exists()
    window.fields["First Name"].setText("Grace")
    assert window.current_qr is None
    assert not window.copy_button.isEnabled()
    assert window.preview.image is None
    window.fields["Email"].setText("broken")
    window.generate_qr()
    assert not window.errors["Email"].isHidden()
    assert window.current_qr is None
    window.close()


def test_batch_does_not_replace_single_contact(qapp, tmp_path):
    window = VCardQRApp()
    window.fields["First Name"].setText("Original")
    window.generate_qr()
    image, payload = window.current_qr, window.current_payload
    path = tmp_path / "contacts.xlsx"
    book = Workbook()
    book.active.append(["Nome"])
    book.active.append(["Batch contact"])
    book.save(path)
    book.close()
    window.load_workbook_path(str(path))
    wait_for_job(qapp, window)
    assert window.mapping_table.cellWidget(0, 1).currentText() == "First Name"
    window.export_to(str(tmp_path), ["First Name"])
    wait_for_job(qapp, window)
    assert "1 saved" in window.batch_summary.text()
    assert window.current_qr is image
    assert window.current_payload == payload
    assert window.save_button.isEnabled()
    window.close()


def test_worker_error_restores_controls(qapp, tmp_path):
    window = VCardQRApp()
    window.load_workbook_path(str(tmp_path / "missing.xlsx"))
    wait_for_job(qapp, window)
    assert window.load_button.isEnabled()
    assert not window.batch_button.isEnabled()
    assert "FileNotFoundError" in window.batch_summary.text()
    window.close()


def test_close_waits_for_cooperative_worker(qapp):
    window = VCardQRApp()
    window.show()

    def operation(task):
        while not task.isInterruptionRequested():
            time.sleep(0.005)
        return {"folder": "test", "saved": 0, "skipped": 0, "failed": 0, "cancelled": True}

    window.begin_job("export", operation)
    window.close()
    assert window.closing
    wait_for_job(qapp, window)
    qapp.processEvents()
    assert not window.isVisible()


def test_custom_remove_theme_and_transparency(qapp):
    window = VCardQRApp()
    window.fields["First Name"].setText("Ada")
    window.add_custom_field()
    key, value = window.custom_fields[0]
    key.setText("Customer ID")
    value.setText("42")
    window.theme.setCurrentText("Dark")
    window.transparent.setChecked(True)
    window.generate_qr()
    assert window.current_qr.mode == "RGBA"
    assert "X-CUSTOMER-ID:42" in window.current_payload
    from PySide6.QtWidgets import QPushButton

    remove = key.parentWidget().findChild(QPushButton)
    remove.click()
    assert not window.custom_fields
    assert window.current_qr is None
    window.close()


def test_denied_clipboard_is_not_reported_as_success(qapp, monkeypatch):
    from PySide6.QtGui import QImage
    from PySide6.QtWidgets import QApplication

    class DeniedClipboard:
        """Model a native session where clipboard writes silently fail."""

        def setImage(self, image):
            pass

        def ownsClipboard(self):
            return False

        def image(self):
            return QImage()

    window = VCardQRApp()
    window.fields["First Name"].setText("Ada")
    window.generate_qr()
    monkeypatch.setattr(QApplication, "clipboard", lambda: DeniedClipboard())
    window.copy_to_clipboard()
    assert "Clipboard unavailable" in window.status.text()
    assert window.save_button.isEnabled()
    window.close()
