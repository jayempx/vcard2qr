"""Inspect spreadsheets and stream row-by-row exports with cancellation."""

import json
import re
from collections.abc import Callable, Sequence
from datetime import datetime
from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook

from .qr import QRSettings, render_qr, save_png
from .vcard import FIELDS, PHONE_TYPES, build_vcard, custom_name

# Preserve the original English/Italian aliases and add common spelling variants.
ALIASES = {re.sub(r"[^a-z0-9]", "", field.lower()): field for field in FIELDS}
ALIASES.update(
    {
        "nome": "First Name",
        "cognome": "Last Name",
        "azienda": "Organization",
        "company": "Organization",
        "cell": "Mobile",
        "cellphone": "Mobile",
        "cellulare": "Mobile",
        "centralino": "Switchboard",
        "office": "Direct Office",
        "directofficephone": "Direct Office",
        "indirizzo": "Address",
        "ruolo": "Title",
    }
)


def cell_text(value: object) -> str:
    """Avoid adding .0 to integral numbers; never invent missing leading zeros."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def default_mapping(headers: Sequence[str]) -> list[str]:
    """Suggest standard fields and explicit custom mappings for unknown headers."""
    return [
        ALIASES.get(re.sub(r"[^a-z0-9]", "", header.lower()), f"Custom: {header}") if header else "Ignore"
        for header in headers
    ]


def validate_mapping(mapping: Sequence[str]) -> None:
    """Reject collisions instead of silently discarding a duplicate column."""
    seen = set()
    for target in mapping:
        if target == "Ignore":
            continue
        name = custom_name(target[8:]) if target.startswith("Custom: ") else target
        if not target.startswith("Custom: ") and target not in FIELDS:
            raise ValueError("Choose a supported destination for every column.")
        if name in seen:
            raise ValueError(f"More than one column maps to {name}. Change one to Ignore.")
        seen.add(name)
    if not seen:
        raise ValueError("Map at least one column before exporting.")


def inspect_workbook(path: str | Path) -> dict:
    """Read only the active sheet's header and first five rows for the preview."""
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = workbook.active
        if sheet is None or sheet.max_column > 200:
            raise ValueError("Use a worksheet with at most 200 columns.")
        rows = sheet.iter_rows()
        header_cells = next(rows, ())
        if any(cell.data_type == "f" for cell in header_cells):
            raise ValueError("Column headers must be text, not formulas.")
        headers = [cell_text(cell.value) for cell in header_cells]
        if not any(headers):
            raise ValueError("Put column headers on the first row of the active sheet.")
        preview = []
        for _ in range(5):
            row = next(rows, None)
            if row is None:
                break
            preview.append([cell_text(cell.value) for cell in row])
        return {
            "headers": headers,
            "preview": preview,
            "sheet": sheet.title,
            "rows": max(0, sheet.max_row - 1),
        }
    finally:
        workbook.close()


def export_workbook(
    path: str | Path,
    output: str | Path,
    headers: Sequence[str],
    mapping: Sequence[str],
    settings: QRSettings,
    cancelled: Callable[[], bool] = lambda: False,
    progress: Callable[[int, int], None] = lambda done, total: None,
) -> dict:
    """Create an isolated export folder and a streaming JSONL per-row report."""
    validate_mapping(mapping)
    if len(headers) != len(mapping):
        raise ValueError("Column mapping no longer matches the workbook. Load it again.")
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows()
        actual = [cell_text(cell.value) for cell in next(rows, ())]
        if actual != list(headers):
            raise ValueError("The workbook headers changed. Load it again before exporting.")
        folder = Path(output) / f"vcard2qr-{datetime.now():%Y%m%d-%H%M%S}-{uuid4().hex[:8]}"
        folder.mkdir(parents=False, exist_ok=False)
        result = {"folder": str(folder), "saved": 0, "skipped": 0, "failed": 0, "cancelled": False}
        total = max(0, sheet.max_row - 1)
        with (folder / "report.jsonl").open("x", encoding="utf-8") as report:
            for row_number, cells in enumerate(rows, start=2):
                if cancelled():
                    result["cancelled"] = True
                    break
                record = {"row": row_number}
                try:
                    data, custom, warnings = {}, [], []
                    for index, target in enumerate(mapping):
                        if target == "Ignore":
                            continue
                        cell = cells[index] if index < len(cells) else None
                        if cell is not None and cell.data_type in ("f", "e"):
                            raise ValueError(
                                f"Column {index + 1}: replace formulas/errors with literal values."
                            )
                        value = cell_text(cell.value if cell else None)
                        if target in PHONE_TYPES and cell and isinstance(cell.value, (int, float)):
                            warnings.append(
                                f"Column {index + 1}: numeric phone; leading zeros cannot be recovered."
                            )
                        if target.startswith("Custom: "):
                            if value:
                                custom.append((target[8:], value))
                        else:
                            data[target] = value
                    if not any(data.values()) and not custom:
                        record["status"] = "skipped"
                        result["skipped"] += 1
                    else:
                        image = render_qr(build_vcard(data, custom), settings)
                        label = " ".join(filter(None, (data.get("First Name"), data.get("Last Name"))))
                        label = label or data.get("Organization") or "contact"
                        # Prefixing the row avoids reserved Windows filenames and duplicates.
                        safe = re.sub(r"[^\w-]+", "_", label, flags=re.UNICODE).strip("_")[:64] or "contact"
                        filename = f"{row_number - 1:04d}_{safe}.png"
                        save_png(image, folder / filename)
                        result["saved"] += 1
                        record.update(status="saved", file=filename, warnings=warnings)
                except (ValueError, OSError) as exc:
                    result["failed"] += 1
                    record.update(status="failed", error=str(exc))
                report.write(json.dumps(record, ensure_ascii=False) + "\n")
                report.flush()
                progress(row_number - 1, total)
        return result
    finally:
        workbook.close()
